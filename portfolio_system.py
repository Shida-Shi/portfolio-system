
import pandas as pd
from pathlib import Path
import shutil
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import DataPoint

# 設定：ファイルパス
BASE_DIR = Path(r"C:\Users\sae\Desktop\python_lesson")
DATA_PATH = BASE_DIR / "蓄積データ.xlsx"
MASTER_PATH = BASE_DIR / "portfolio - コピー.xlsx"
TEMPLATE_PATH = BASE_DIR / "入力テンプレート.xlsx"

# -----------------------------------------------
# 0. マスタデータの読み込み
# -----------------------------------------------
def マスタ読み込み():
    try:
        df_ms = pd.read_excel(MASTER_PATH, sheet_name="マスタ", dtype=str)
        return df_ms
    except Exception as e:
        print(f"マスタの読み込みに失敗しました: {e}")
        return None

# -----------------------------------------------
# 1. データ登録（1件）
# -----------------------------------------------
def データ登録(月, 会社名, 商品名, 店番, 税抜金額):
    df_ms = マスタ読み込み()
    グループ名 = "未登録"
    店舗名 = "未登録"

    if df_ms is not None:
        match = df_ms[df_ms["店舗コード"].astype(str) == str(店番)]
        if not match.empty:
            グループ名 = match.iloc[0]["グループ名"]
            店舗名 = match.iloc[0]["店舗名"]
            会社名 = match.iloc[0]["会社名"]

    if DATA_PATH.exists():
        df = pd.read_excel(DATA_PATH, dtype={"店番": str})
    else:
        df = pd.DataFrame(columns=["月","会社名","グループ名","店番","店舗名","商品名","税抜金額","消費税","税込金額"])

    消費税 = int(税抜金額 * 0.1)
    税込金額 = int(税抜金額 * 1.1)
    新しい行 = {
        "月": 月, "会社名": 会社名, "グループ名": グループ名,
        "店番": 店番, "店舗名": 店舗名, "商品名": 商品名,
        "税抜金額": 税抜金額, "消費税": 消費税, "税込金額": 税込金額
    }
    df = pd.concat([df, pd.DataFrame([新しい行])], ignore_index=True)
    df.to_excel(DATA_PATH, index=False)

    print("-" * 30)
    print(f"【登録完了】")
    print(f"店舗: {店舗名} ({グループ名})")
    print(f"内容: {会社名} / {商品名}")
    print(f"金額: {税込金額:,}円 (内消費税 {消費税:,}円)")
    print("-" * 30)

# -----------------------------------------------
# 2. テンプレート生成
# -----------------------------------------------
def テンプレート生成():
    df = pd.DataFrame(columns=["月", "店番", "税抜金額"])
    df.to_excel(TEMPLATE_PATH, index=False)
    print(f"テンプレートを作成しました → {TEMPLATE_PATH}")
    print("Excelで開いて月・店番・税抜金額を入力してください。")

# -----------------------------------------------
# 3. 一括登録（Excelテンプレートから）
# -----------------------------------------------
def 一括登録():
    if not TEMPLATE_PATH.exists():
        print("テンプレートが見つかりません。先に「2. テンプレート生成」を実行してください。")
        return
    try:
        df_input = pd.read_excel(TEMPLATE_PATH, dtype={"店番": str})
    except Exception as e:
        print(f"テンプレートの読み込みに失敗しました: {e}")
        return

    df_input = df_input.dropna(subset=["月", "店番", "税抜金額"])
    if df_input.empty:
        print("テンプレートにデータがありません。")
        return

    df_ms = マスタ読み込み()

    if DATA_PATH.exists():
        df = pd.read_excel(DATA_PATH, dtype={"店番": str})
    else:
        df = pd.DataFrame(columns=["月","会社名","グループ名","店番","店舗名","商品名","税抜金額","消費税","税込金額"])

    成功 = 0
    スキップ = 0
    スキップ一覧 = []

    for _, row in df_input.iterrows():
        月 = int(row["月"])
        店番 = str(row["店番"]).strip()
        税抜金額 = int(row["税抜金額"])

        グループ名 = 会社名 = 店舗名 = 商品名 = "未登録"

        if df_ms is not None:
            match = df_ms[df_ms["店舗コード"].astype(str) == 店番]
            if not match.empty:
                グループ名 = match.iloc[0]["グループ名"]
                店舗名 = match.iloc[0]["店舗名"]
                会社名 = match.iloc[0]["会社名"]
                商品名 = match.iloc[0]["商品名"]
            else:
                スキップ += 1
                スキップ一覧.append(店番)
                continue

        消費税 = int(税抜金額 * 0.1)
        税込金額 = int(税抜金額 * 1.1)
        新しい行 = {
            "月": 月, "会社名": 会社名, "グループ名": グループ名,
            "店番": 店番, "店舗名": 店舗名, "商品名": 商品名,
            "税抜金額": 税抜金額, "消費税": 消費税, "税込金額": 税込金額
        }
        df = pd.concat([df, pd.DataFrame([新しい行])], ignore_index=True)
        成功 += 1

    if 成功 > 0:
        df.to_excel(DATA_PATH, index=False)

    print("-" * 30)
    print(f"【一括登録完了】")
    print(f"✅ 登録成功: {成功}件")
    if スキップ > 0:
        print(f"⚠ スキップ（マスタ未登録）: {スキップ}件")
        print(f"  店番: {', '.join(スキップ一覧)}")
    print("-" * 30)

# -----------------------------------------------
# 4. 集計
# -----------------------------------------------
def 集計(会社名, 商品名, 按分率=1.0):
    if not DATA_PATH.exists():
        print("データがまだありません")
        return
    df = pd.read_excel(DATA_PATH)
    対象 = df[(df["会社名"] == 会社名) & (df["商品名"] == 商品名)]
    if 対象.empty:
        print("該当データがありません")
        return
    月別 = 対象.groupby("月")["税抜金額"].sum()
    合計 = 月別.sum()
    按分 = int(合計 * 按分率)
    print(f"\n--- {会社名} / {商品名} ---")
    print(月別.to_string())
    print(f"合計: {合計:,}円")
    print(f"按分({int(按分率*100)}%): {按分:,}円")

# -----------------------------------------------
# 5. 一括出力（グラフ付き）
# -----------------------------------------------
def 一括出力(按分率=1.0):
    if not DATA_PATH.exists():
        print("データがまだありません")
        return
    df = pd.read_excel(DATA_PATH)
    出力フォルダ = BASE_DIR / f"一括出力_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    出力フォルダ.mkdir()
    組み合わせ = df[["会社名","商品名"]].drop_duplicates()
    count = 0

    for _, row in 組み合わせ.iterrows():
        会社名 = row["会社名"]
        商品名 = row["商品名"]
        対象 = df[(df["会社名"] == 会社名) & (df["商品名"] == 商品名)]

        # 月別集計（1〜12月を全部並べる）
        月別 = 対象.groupby("月")["税抜金額"].sum().reindex(range(1, 13), fill_value=0).reset_index()
        月別.columns = ["月", "税抜金額"]
        月別["按分金額"] = (月別["税抜金額"] * 按分率).astype(int)

        ファイル名 = f"{会社名}_{商品名}.xlsx"
        出力パス = 出力フォルダ / ファイル名
        月別.to_excel(出力パス, index=False, sheet_name="月別集計")

        # グラフをExcelに追加
        wb = load_workbook(出力パス)
        ws = wb["月別集計"]

        chart = BarChart()
        chart.type = "col"
        chart.title = f"{会社名} / {商品名} 月別売上"
        chart.y_axis.title = "税抜金額（円）"
        chart.x_axis.title = "月"
        chart.style = 10
        chart.width = 20
        chart.height = 12

        # データ範囲（税抜金額の列）
        data = Reference(ws, min_col=2, min_row=1, max_row=13)
        cats = Reference(ws, min_col=1, min_row=2, max_row=13)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)

        ws.add_chart(chart, "E2")
        wb.save(出力パス)

        count += 1
        print(f"出力: {ファイル名}")

    print(f"\n✅ {count}件を出力しました（グラフ付き）→ {出力フォルダ}")

# -----------------------------------------------
# 6. バックアップ
# -----------------------------------------------
def バックアップ():
    if not DATA_PATH.exists():
        print("バックアップするデータがありません")
        return
    now = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = BASE_DIR / f"バックアップ_{now}.xlsx"
    shutil.copy2(DATA_PATH, backup_path)
    print(f"バックアップ完了: {backup_path.name}")

# -----------------------------------------------
# 7. 年度更新
# -----------------------------------------------
def 年度更新():
    確認 = input("【重要】全データを消去して年度更新します。よろしいですか？(yes/no): ")
    if 確認 != "yes":
        print("キャンセルしました")
        return
    バックアップ()
    df = pd.DataFrame(columns=["月","会社名","グループ名","店番","店舗名","商品名","税抜金額","消費税","税込金額"])
    df.to_excel(DATA_PATH, index=False)
    print("年度更新完了！データをリセットしました。")

# -----------------------------------------------
# メニュー
# -----------------------------------------------
def メニュー():
    while True:
        print("\n=============================")
        print("  ポートフォリオ集計システム")
        print("=============================")
        print("1. データ登録（1件）")
        print("2. テンプレート生成")
        print("3. 一括登録（Excelから）")
        print("4. 集計表示")
        print("5. 一括出力（グラフ付き）")
        print("6. バックアップ")
        print("7. 年度更新")
        print("0. 終了")
        選択 = input("番号を入力してください: ")

        if 選択 == "1":
            try:
                月 = int(input("月 (1-12): "))
                店番 = input("店番を入力してください: ")
                df_ms = マスタ読み込み()
                会社名 = ""
                商品名 = ""
                if df_ms is not None:
                    match = df_ms[df_ms["店舗コード"].astype(str) == str(店番)]
                    if not match.empty:
                        print(f" -> 判定: {match.iloc[0]['会社名']} / {match.iloc[0]['店舗名']}")
                        商品名 = match.iloc[0]["商品名"]
                        入力商品 = input(f"商品名 (空欄なら '{商品名}'): ")
                        if 入力商品 != "":
                            商品名 = 入力商品
                        会社名 = match.iloc[0]["会社名"]
                    else:
                        print(" -> マスタにない店番です。手動入力してください。")
                        会社名 = input("会社名: ")
                        商品名 = input("商品名: ")
                税抜金額 = int(input("税抜金額: "))
                データ登録(月, 会社名, 商品名, 店番, 税抜金額)
            except ValueError:
                print("【エラー】数値の入力が正しくありません。")

        elif 選択 == "2":
            テンプレート生成()

        elif 選択 == "3":
            一括登録()

        elif 選択 == "4":
            会社名 = input("会社名: ")
            商品名 = input("商品名: ")
            按分率 = float(input("按分率 (例: 0.5): "))
            集計(会社名, 商品名, 按分率)

        elif 選択 == "5":
            按分率 = float(input("按分率 (例: 1.0): "))
            一括出力(按分率)

        elif 選択 == "6":
            バックアップ()

        elif 選択 == "7":
            年度更新()

        elif 選択 == "0":
            print("終了します")
            break

if __name__ == "__main__":
    メニュー()